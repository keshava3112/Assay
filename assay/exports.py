"""Exports keep the receipts: every cell has its status, quote and source URL."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from .domain import Run, SetSpec, to_jsonable


def _flat(spec: SetSpec, run: Run) -> tuple[list[str], list[list[str]]]:
    header = ["company", "website", "match", "score"]
    for c in spec.criteria:
        header += [c.label, f"{c.label} (quote)", f"{c.label} (source)"]
    for e in spec.enrich:
        header += [e.label, f"{e.label} (source)"]
    rows = []
    for r in run.rows:
        line = [r.name, r.homepage, "yes" if r.is_match(spec) else "no", f"{r.score(spec):.2f}"]
        for c in spec.criteria:
            v = r.verdicts.get(c.id)
            ev = v.evidence[0] if v and v.evidence else None
            line += [v.status.value if v else "", ev.quote if ev else (v.rationale if v else ""), ev.url if ev else ""]
        for e in spec.enrich:
            f = r.fields.get(e.id)
            ev = f.evidence[0] if f and f.evidence else None
            line += [(f.value or "") if f else "", ev.url if ev else ""]
        rows.append(line)
    return header, rows


def to_csv(spec: SetSpec, run: Run, path: str | Path) -> Path:
    header, rows = _flat(spec, run)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    return Path(path)


def to_json(spec: SetSpec, run: Run, path: str | Path) -> Path:
    Path(path).write_text(json.dumps({"set": to_jsonable(spec), "run": to_jsonable(run)}, indent=1, ensure_ascii=False))
    return Path(path)


def to_xlsx(spec: SetSpec, run: Run, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header, rows = _flat(spec, run)
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="F5E7B8"), PatternFill("solid", fgColor="141416")
    fills = {"pass": "E3F0E6", "fail": "F6E1DC", "unclear": "EFEDE6"}
    for r in rows:
        ws.append(r)
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.value in fills:
                cell.fill = PatternFill("solid", fgColor=fills[cell.value])
    for col, width in zip("ABCD", (22, 26, 8, 8)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B2"
    wb.save(path)
    return Path(path)
