import asyncio
import csv
import json
from datetime import datetime

import pytest
from conftest import AUG, OCT, make_engine, run
from openpyxl import load_workbook
from starlette.testclient import TestClient

from assay.api import create_app
from assay.config import SpecError, parse_spec
from assay.domain import Status
from assay.exports import to_csv, to_json, to_xlsx
from assay.service import execute
from assay.site import build_site
from assay.store import Store


@pytest.fixture()
def store(tmp_path, raw_spec, leads):
    s = Store(tmp_path / "assay.db")
    s.put_set(raw_spec, leads)
    return s


@pytest.fixture()
def seeded(store, spec, leads):
    run(execute(store, spec.id, make_engine("aug", leads, AUG)))
    run(execute(store, spec.id, make_engine("oct", leads, OCT)))
    return store


def test_spec_validation(raw_spec):
    with pytest.raises(SpecError):
        parse_spec({**raw_spec, "criteria": [{"id": "x", "signals": ["(unclosed"]}]})
    with pytest.raises(SpecError):
        parse_spec({k: v for k, v in raw_spec.items() if k != "query"})


def test_store_round_trip(seeded, spec):
    runs = seeded.runs_for(spec.id)
    assert len(runs) == 2
    latest = seeded.latest_run(spec.id)
    again = seeded.get_run(latest.id)
    assert again.rows[0].verdicts.keys() == latest.rows[0].verdicts.keys()
    assert {r.name for r in again.rows if r.is_match(spec)} == {"Kosha Credit", "Mitra Money", "Saral Loans", "Udhaar Grid"}
    assert latest.changes and any(c.name == "Udhaar Grid" for c in latest.changes)
    ev = seeded.events(latest.id)
    assert ev[0]["kind"] == "run.started" and ev[-1]["kind"] == "run.saved"
    assert [e["seq"] for e in ev] == list(range(len(ev)))


def test_exports(seeded, spec, tmp_path):
    r = seeded.latest_run(spec.id)
    rows = list(csv.reader(open(to_csv(spec, r, tmp_path / "x.csv"))))
    assert rows[0][:5] == ["company", "website", "match", "score", "HQ in India"]
    assert rows[0][5:7] == ["HQ in India (quote)", "HQ in India (source)"]
    kosha = next(x for x in rows if x[0] == "Kosha Credit")
    assert kosha[2] == "yes" and "headquartered in Bengaluru" in kosha[5]
    assert len(rows) == len(r.rows) + 1
    data = json.loads(to_json(spec, r, tmp_path / "x.json").read_text())
    assert data and isinstance(data, (list, dict))
    wb = load_workbook(to_xlsx(spec, r, tmp_path / "x.xlsx"))
    assert wb.active.max_row == len(r.rows) + 1


def test_api_reads(seeded, spec):
    c = TestClient(create_app(seeded))
    assert c.get("/healthz").json() == {"ok": True}
    boot = c.get("/api/bootstrap").json()
    assert boot["mode"] == "live" and boot["sets"][0]["id"] == spec.id
    d = c.get(f"/api/sets/{spec.id}").json()
    assert d["run"]["id"] and len(d["history"]) == 2
    assert c.get("/api/sets/nope").status_code == 404
    x = c.get(f"/api/sets/{spec.id}/export.csv")
    assert x.status_code == 200 and x.headers["content-type"].startswith("text/csv")
    assert c.get(f"/api/sets/{spec.id}/export.exe").status_code == 404
    assert "Assay" in c.get("/").text


def test_api_token(seeded, monkeypatch):
    monkeypatch.setenv("ASSAY_TOKEN", "s3cret")
    c = TestClient(create_app(seeded))
    assert c.get("/api/bootstrap").status_code == 401
    assert c.get("/api/bootstrap", headers={"authorization": "Bearer s3cret"}).status_code == 200
    assert c.get("/healthz").status_code == 200


class SlowEngine:
    """Wraps a real engine but waits for a gate, so a second POST lands mid-run."""

    def __init__(self, inner, gate):
        self.inner, self.gate = inner, gate

    async def run(self, spec, emit=None, run_id=None):
        await self.gate.wait()
        return await self.inner.run(spec, emit=emit, run_id=run_id)


def test_api_run_streams_events_and_refuses_overlap(store, spec, leads):
    gate = asyncio.Event()
    app = create_app(store, engine_factory=lambda s, l: SlowEngine(make_engine("oct", leads, OCT), gate))
    with TestClient(app) as c:
        r = c.post(f"/api/sets/{spec.id}/runs")
        assert r.status_code == 202
        run_id = r.json()["run_id"]
        assert c.post(f"/api/sets/{spec.id}/runs").status_code == 409
        c.portal.call(gate.set)
        kinds = []
        with c.stream("GET", f"/api/runs/{run_id}/events") as s:
            for line in s.iter_lines():
                if line.startswith("event: "):
                    kinds.append(line[7:])
        assert kinds[0] == "run.started" and kinds[-1] == "run.saved" and "verdict" in kinds
        got = c.get(f"/api/runs/{run_id}").json()
        assert got["id"] == run_id and got["stats"]["matches"] == 4


def test_static_site_bakes_real_data(seeded, spec, tmp_path):
    out = build_site(seeded, tmp_path / "dist")
    html = (out / "index.html").read_text()
    assert 'src="static/data.js"' in html and 'href="/static/' not in html and 'src="/static/' not in html
    assert not (out / "static" / "index.html").exists() and (out / ".nojekyll").exists()
    js = (out / "static" / "data.js").read_text()
    data = json.loads(js.removeprefix("window.ASSAY_DATA = ").rstrip().rstrip(";"))
    assert data["mode"] == "static" and data["sets"][0]["id"] == spec.id
    run_id = data["details"][spec.id]["run"]["id"]
    assert data["events"][run_id][-1]["kind"] == "run.saved"
    for name in ("app.js", "styles.css", "logo.svg"):
        assert (out / "static" / name).exists()


def test_gate_downgrades_are_counted(store, spec, leads):
    class LyingJudge:
        async def judge(self, spec, candidate, criterion, passages, now):
            from assay.domain import Evidence, Verdict

            p = candidate.pages[0]
            return Verdict(criterion.id, Status.PASS, "invented",
                           [Evidence(p.url, "This sentence is not on the page at all.", p.content_hash, p.fetched_at)], 0.99)

        async def enrich(self, spec, candidate, enrichment, passages):
            from assay.domain import Field

            return Field(enrichment.id, None)

    r = run(make_engine("oct", leads, OCT, judge=LyingJudge()).run(spec))
    assert r.stats.matches == 0 and r.stats.gated > 0
    assert all(v.status is Status.UNCLEAR for row in r.rows for v in row.verdicts.values())
